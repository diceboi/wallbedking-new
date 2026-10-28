import openpyxl
import json
import sys

wb = openpyxl.load_workbook('New Products 2026.xlsx', data_only=True)

def clean_val(v):
    if v is None:
        return None
    return str(v).strip()

for name in wb.sheetnames:
    ws = wb[name]
    print(f"\n==================== SHEET: {name} ====================")
    all_rows = list(ws.iter_rows(values_only=True))
    non_empty = []
    for idx, r in enumerate(all_rows):
        cleaned = [clean_val(c) for c in r]
        if any(c is not None and c != "" for c in cleaned):
            non_empty.append((idx + 1, cleaned))
    
    print(f"Total non-empty rows: {len(non_empty)}")
    for row_num, row_data in non_empty[:15]:
        filtered = [(col_idx, val) for col_idx, val in enumerate(row_data) if val is not None and val != ""]
        print(f"Row {row_num}: {filtered[:8]}")
