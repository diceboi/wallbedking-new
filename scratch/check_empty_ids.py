import openpyxl

wb = openpyxl.load_workbook(r'C:\Users\szasz\Downloads\wallbedking-products-filtered-2026-10-02%20%282%29.xlsx', data_only=True)
ws = wb.active
raw_headers = [cell.value for cell in ws[1]]

for r_idx in range(2, ws.max_row + 1):
    vals = [cell.value for cell in ws[r_idx]]
    if not any(vals): continue
    d = {h: v for h, v in zip(raw_headers, vals) if h}
    if d.get('ID') is None or str(d.get('ID')).strip() == '':
        print(f"Row {r_idx}: SKU={d.get('SKU')}, Name={d.get('Product Name (Default)')}")
        for k, v in d.items():
            if v is not None:
                print(f"   {k}: {repr(v)[:100]}")
