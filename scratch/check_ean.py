import openpyxl, json, urllib.request

supabase_url = 'https://unrqbejocbteebsworuq.supabase.co'
service_key = 'eyJhbGciOiJIUzI1NiIsInR5cCI6IkpXVCJ9.eyJpc3MiOiJzdXBhYmFzZSIsInJlZiI6InVucnFiZWpvY2J0ZWVic3dvcnVxIiwicm9sZSI6InNlcnZpY2Vfcm9sZSIsImlhdCI6MTc4NzYzNTU4NywiZXhwIjoyMTAzMjExNTg3fQ.wMOHLiuny-ZRwTL3Fw03H_bZk20FbEIjaj0WGZlcIn4'

req = urllib.request.Request(
    f'{supabase_url}/rest/v1/products?select=id,ean,gtin_en&limit=5',
    headers={'apikey': service_key, 'Authorization': f'Bearer {service_key}'}
)
with urllib.request.urlopen(req) as resp:
    db_products = json.loads(resp.read().decode('utf-8'))

wb = openpyxl.load_workbook(r'C:\Users\szasz\Downloads\wallbedking-products-filtered-2026-10-02%20%282%29.xlsx', data_only=True)
ws = wb.active
for p in db_products:
    pid = str(p['id'])
    for row in range(2, 8):
        if str(ws.cell(row=row, column=1).value) == pid:
            print(f"ID {pid}: DB ean={repr(p.get('ean'))}, gtin_en={repr(p.get('gtin_en'))} | Excel EAN={repr(ws.cell(row=row, column=54).value)}")
