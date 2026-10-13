import openpyxl, sys

wb = openpyxl.load_workbook(r'C:\Users\szasz\Downloads\wallbedking-products-filtered-2026-10-02%20%282%29.xlsx', data_only=True)
ws = wb.active
headers = [c.value for c in ws[1]]

for r in ws.iter_rows(min_row=2, values_only=True):
    d = {h: v for h, v in zip(headers, r) if h}
    if str(d.get('ID')) in ['52', '53']:
        sys.stdout.buffer.write(f"ID {d.get('ID')}:\n".encode('utf-8'))
        sys.stdout.buffer.write(f"  Default: {d.get('Product Name (Default)')}\n".encode('utf-8'))
        sys.stdout.buffer.write(f"  EN:      {d.get('Name EN (UK)')}\n".encode('utf-8'))
        sys.stdout.buffer.write(f"  DE:      {d.get('Name DE')}\n".encode('utf-8'))
        sys.stdout.buffer.write(f"  FR:      {d.get('Name FR')}\n".encode('utf-8'))
        sys.stdout.buffer.write(f"  Dim:     {d.get('Dimensions (cm)')}\n".encode('utf-8'))
